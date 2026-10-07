import io
import json
import logging
import os
import re
import sqlite3
import unicodedata
import zipfile
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Literal

from docx import Document
from lxml.etree import XMLSyntaxError
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader
from pypdf.errors import PdfReadError

logger = logging.getLogger(__name__)
DB_PATH = Path(os.environ.get("CV_REVIEW_DB", Path(__file__).parent / "data" / "reviews.db"))
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_LENGTH = 150_000


@contextmanager
def database():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


class JobInput(BaseModel):
    title: str = Field(min_length=2, max_length=120)
    department: str = Field(min_length=2, max_length=80)
    description: str = Field(min_length=20, max_length=20_000)
    skills: list[str] = Field(min_length=1, max_length=30)
    min_years: float = Field(ge=0, le=50)
    skill_weight: int = Field(default=80, ge=0, le=100)

    @field_validator("title", "department", "description", mode="before")
    @classmethod
    def trim_string(cls, value: object) -> str:
        if not isinstance(value, str):
            raise ValueError("Trường này phải là văn bản.")
        return value.strip()

    @field_validator("skills")
    @classmethod
    def clean_skills(cls, values: list[str]) -> list[str]:
        cleaned = list(dict.fromkeys(value.strip() for value in values))
        if any(not value or len(value) > 60 for value in cleaned):
            raise ValueError("Kỹ năng phải có từ 1 đến 60 ký tự.")
        return cleaned


class ReviewInput(BaseModel):
    status: Literal["pending", "shortlisted", "hold", "not_fit"]
    notes: str = Field(default="", max_length=5000)


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.casefold().replace("đ", "d"))
    return "".join(char for char in decomposed if not unicodedata.combining(char))


ALIASES = {
    "react": ["react", "reactjs", "react.js"],
    "node.js": ["node.js", "nodejs", "node js"],
    "javascript": ["javascript", "js"],
    "typescript": ["typescript", "ts"],
    "postgresql": ["postgresql", "postgres"],
    "python": ["python"],
}
SKILL_TERMS = {normalize(alias): aliases for aliases in ALIASES.values() for alias in aliases}


def evaluate(text: str, job: dict) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    evidence = []
    for skill in job["skills"]:
        terms = SKILL_TERMS.get(normalize(skill), [skill])
        pattern = re.compile(
            r"(?<![\w])(?:" + "|".join(re.escape(normalize(term)) for term in terms) + r")(?![\w])"
        )
        matches = [line[:350] for line in lines if pattern.search(normalize(line))]
        evidence.append({"skill": skill, "matched": bool(matches), "evidence": matches[:3]})

    # Only explicit experience claims count; calendar dates are not inferred.
    years_matches = []
    experience_pattern = re.compile(
        r"(\d+(?:[.,]\d+)?)\s*\+?\s*(?:years?|nam)\s*"
        r"(?:(?:of|kinh)\s*)?(?:experience|kinh nghiem|nghiem)",
    )
    for line in lines:
        for match in experience_pattern.finditer(normalize(line)):
            value = float(match.group(1).replace(",", "."))
            if value <= 50:
                years_matches.append((value, line[:350]))
    years = max((value for value, _ in years_matches), default=None)
    matched = sum(item["matched"] for item in evidence)
    skill_score = matched / len(evidence) * 100
    experience_score = (
        100 if job["min_years"] == 0
        else min(100, years / job["min_years"] * 100) if years is not None
        else 0
    )
    score = round(
        skill_score * job["skill_weight"] / 100
        + experience_score * (100 - job["skill_weight"]) / 100
    )
    return {
        "score": score,
        "skill_score": round(skill_score),
        "experience_score": round(experience_score),
        "matched_skills": matched,
        "total_skills": len(evidence),
        "skills": evidence,
        "years": years,
        "experience_evidence": [line for value, line in years_matches if value == years][:3],
        "warnings": (
            ["Chưa tìm thấy số năm kinh nghiệm được khai báo rõ; cần kiểm tra thủ công."]
            if years is None and job["min_years"] > 0 else []
        ),
    }


def extract_text(content: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise HTTPException(422, "PDF có mật khẩu. Vui lòng tải bản không mã hóa.")
            if len(reader.pages) > 100:
                raise HTTPException(422, "CV không được vượt quá 100 trang.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(422, "Nội dung DOCX giải nén quá lớn.")
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = "\n".join(parts)
        else:
            raise HTTPException(415, "Chỉ hỗ trợ PDF, DOCX và TXT (UTF-8).")
    except (UnicodeDecodeError, PdfReadError, zipfile.BadZipFile, XMLSyntaxError, KeyError, ValueError) as error:
        logger.warning("Cannot parse uploaded CV: %s", type(error).__name__)
        raise HTTPException(422, "Không đọc được tệp. Kiểm tra định dạng và nội dung CV.") from error
    if len(text.strip()) < 30:
        raise HTTPException(422, "CV không đủ văn bản để đánh giá. PDF scan cần OCR trước khi tải lên.")
    if len(text) > MAX_TEXT_LENGTH:
        raise HTTPException(422, "CV có quá nhiều nội dung (tối đa 150.000 ký tự).")
    return text.strip()


def job_dict(row: sqlite3.Row) -> dict:
    result = dict(row)
    result["skills"] = json.loads(result["skills"])
    return result


def candidate_dict(row: sqlite3.Row) -> dict:
    result = dict(row)
    result["assessment"] = json.loads(result["assessment"])
    return result


def insert_job(connection: sqlite3.Connection, job: JobInput) -> int:
    cursor = connection.execute(
        "INSERT INTO jobs (title,department,description,skills,min_years,skill_weight) VALUES (?,?,?,?,?,?)",
        (job.title, job.department, job.description, json.dumps(job.skills), job.min_years, job.skill_weight),
    )
    assert cursor.lastrowid is not None
    return cursor.lastrowid


@asynccontextmanager
async def lifespan(app: FastAPI):
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with database() as connection:
        connection.executescript("""
            CREATE TABLE IF NOT EXISTS jobs (
                id INTEGER PRIMARY KEY, title TEXT NOT NULL, department TEXT NOT NULL,
                description TEXT NOT NULL, skills TEXT NOT NULL,
                min_years REAL NOT NULL, skill_weight INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS candidates (
                id INTEGER PRIMARY KEY, job_id INTEGER NOT NULL REFERENCES jobs(id),
                name TEXT NOT NULL, email TEXT NOT NULL, filename TEXT NOT NULL,
                text TEXT NOT NULL, assessment TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending', notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );
        """)
        if connection.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 0:
            insert_job(connection, JobInput(
                title="Senior Frontend Developer", department="Engineering",
                description="Phát triển giao diện web với React và TypeScript. Xây dựng UI dễ sử dụng, "
                "tối ưu hiệu năng, phối hợp cùng thiết kế và backend. Đây là JD mẫu, hãy tạo JD của bạn.",
                skills=["React", "TypeScript", "JavaScript", "HTML", "CSS", "Git"],
                min_years=3,
            ))
    yield


app = FastAPI(title="TalentLens API", version="1.0.0", lifespan=lifespan)


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/jobs")
def list_jobs():
    with database() as connection:
        return [job_dict(row) for row in connection.execute("SELECT * FROM jobs ORDER BY id DESC")]


@app.post("/api/jobs", status_code=201)
def create_job(job: JobInput):
    with database() as connection:
        job_id = insert_job(connection, job)
        return job_dict(connection.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone())


@app.get("/api/jobs/{job_id}/candidates")
def list_candidates(job_id: int):
    with database() as connection:
        if not connection.execute("SELECT id FROM jobs WHERE id=?", (job_id,)).fetchone():
            raise HTTPException(404, "Không tìm thấy JD.")
        return [candidate_dict(row) for row in connection.execute(
            "SELECT * FROM candidates WHERE job_id=? ORDER BY id DESC", (job_id,)
        )]


@app.post("/api/jobs/{job_id}/candidates", status_code=201)
async def upload_candidate(job_id: int, file: UploadFile = File(...)):
    with database() as connection:
        row = connection.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "Không tìm thấy JD.")
        job = job_dict(row)
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "Mỗi CV tối đa 10 MB.")
    filename = (file.filename or "cv").replace("\\", "/").split("/")[-1][:200]
    text = extract_text(content, filename)
    name = Path(filename).stem.replace("_", " ").replace("-", " ").strip() or "Ứng viên"
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}", text)
    assessment = evaluate(text, job)
    with database() as connection:
        cursor = connection.execute(
            "INSERT INTO candidates (job_id,name,email,filename,text,assessment,created_at) VALUES (?,?,?,?,?,?,?)",
            (job_id, name, email_match.group(0) if email_match else "", filename,
             text, json.dumps(assessment, ensure_ascii=False), datetime.now(timezone.utc).isoformat()),
        )
        return candidate_dict(connection.execute(
            "SELECT * FROM candidates WHERE id=?", (cursor.lastrowid,)
        ).fetchone())


@app.patch("/api/candidates/{candidate_id}")
def update_review(candidate_id: int, review: ReviewInput):
    with database() as connection:
        cursor = connection.execute(
            "UPDATE candidates SET status=?, notes=? WHERE id=?",
            (review.status, review.notes, candidate_id),
        )
        if cursor.rowcount == 0:
            raise HTTPException(404, "Không tìm thấy CV.")
        return candidate_dict(connection.execute(
            "SELECT * FROM candidates WHERE id=?", (candidate_id,)
        ).fetchone())


@app.delete("/api/candidates/{candidate_id}", status_code=204)
def delete_candidate(candidate_id: int):
    with database() as connection:
        if connection.execute("DELETE FROM candidates WHERE id=?", (candidate_id,)).rowcount == 0:
            raise HTTPException(404, "Không tìm thấy CV.")
